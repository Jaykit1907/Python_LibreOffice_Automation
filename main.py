from pathlib import Path
import shutil
import subprocess
import openpyxl
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT

# ---------------------------------------------------------
# CONFIGURATION
# ---------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"

EXCEL_FILE = INPUT_DIR / "Andhra Pradesh.xlsx"
TEMPLATE_FILE = INPUT_DIR / "Template.docx"
WORD_OUTPUT = OUTPUT_DIR / "Andhra_Pradesh_District_Report.docx"
PDF_OUTPUT = OUTPUT_DIR / "Andhra_Pradesh_District_Report.pdf"


# ---------------------------------------------------------
# STEP 1: READ EXCEL DATA
# Each worksheet is a district. A row whose second cell is
# "Damaged" marks the start of a location's product table.
# ---------------------------------------------------------
def read_excel_data(file_path):
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    districts = []

    for sheet in workbook.worksheets:
        locations = []
        current_location = None

        for row in sheet.iter_rows(values_only=True):
            first = row[0] if len(row) > 0 else None
            second = row[1] if len(row) > 1 else None

            if first is None:
                continue

            if str(second).strip().lower() == "damaged":
                current_location = {
                    "name": str(first).strip(),
                    "products": []
                }
                locations.append(current_location)
                continue

            if current_location is not None:
                damaged = row[1] if len(row) > 1 else None
                expired = row[2] if len(row) > 2 else None
                total = row[3] if len(row) > 3 else None

                if all(value is not None for value in
                       (damaged, expired, total)):
                    current_location["products"].append([
                        str(first).strip(),
                        str(damaged),
                        str(expired),
                        str(total)
                    ])

        locations = [loc for loc in locations if loc["products"]]
        if locations:
            districts.append({
                "name": sheet.title,
                "locations": locations
            })

    return districts


# ---------------------------------------------------------
# STEP 2: REPLACE TEMPLATE PLACEHOLDERS
# ---------------------------------------------------------
def replace_placeholder(document, placeholder, value):
    for paragraph in document.paragraphs:
        if placeholder in paragraph.text:
            paragraph.text = paragraph.text.replace(
                placeholder, str(value)
            )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if placeholder in paragraph.text:
                        paragraph.text = paragraph.text.replace(
                            placeholder, str(value)
                        )


# ---------------------------------------------------------
# STEP 3: CREATE LOCATION TABLES
# ---------------------------------------------------------

def format_cell(cell, value, bold=False):
    cell.text = str(value)

    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)

        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.bold = bold


def add_location_table(document, location, index):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.35)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(5)

    run = paragraph.add_run(
        f"Location-{index}: {location['name']}"
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.bold = True

    # Create the table
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Explicit column widths
    widths = [
        Inches(2.4),
        Inches(1.0),
        Inches(1.5),
        Inches(1.0)
    ]

    headers = [
        "Product Name",
        "Damaged",
        "Date of Expired",
        "Total"
    ]

    # Format header cells
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.width = widths[col]
        format_cell(cell, header, bold=True)

    # Format data cells
    for product in location["products"]:
        cells = table.add_row().cells

        for col, value in enumerate(product):
            cells[col].width = widths[col]
            format_cell(cells[col], value)

    # Add visible borders to every cell
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = tc_pr.find(qn("w:tcBorders"))

            if borders is None:
                borders = OxmlElement("w:tcBorders")
                tc_pr.append(borders)

            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                tag = qn(f"w:{edge}")
                element = borders.find(tag)

                if element is None:
                    element = OxmlElement(f"w:{edge}")
                    borders.append(element)

                element.set(qn("w:val"), "single")
                element.set(qn("w:sz"), "6")
                element.set(qn("w:space"), "0")
                element.set(qn("w:color"), "000000")

    document.add_paragraph()

# ---------------------------------------------------------
# STEP 4: GENERATE ONE WORD REPORT PER WORKSHEET
# ---------------------------------------------------------
def generate_word(template_path, output_dir, district):
    # Start with a fresh copy of the template for each worksheet.
    document = Document(template_path)

    district_name = district["name"]
    locations = district["locations"]
    total_locations = len(locations)

    # Replace placeholders using this worksheet's data.
    replace_placeholder(
        document, "<DistrictName>", district_name
    )
    replace_placeholder(
        document, "<TotalLocations>", total_locations
    )

    # Add only this worksheet's district data.
    document.add_page_break()

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(8)

    run = heading.add_run(f"District Name: {district_name}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)

    document.add_paragraph(
        f"Total Locations: {total_locations}"
    )

    for index, location in enumerate(locations, start=1):
        add_location_table(document, location, index)

    # Make the worksheet name safe for Windows filenames.
    safe_name = "".join(
        "_" if char in '<>:"/\\|?*' else char
        for char in district_name
    ).strip().rstrip(".")

    if not safe_name:
        safe_name = "District"

    output_path = output_dir / f"{safe_name}.docx"
    document.save(output_path)

    print(f"Word report saved: {output_path}")
    return output_path


# ---------------------------------------------------------
# STEP 5: CONVERT WORD TO PDF USING LIBREOFFICE
# ---------------------------------------------------------
def convert_to_pdf(word_path, output_directory):
    # First check the known LibreOffice installation path.
    possible_paths = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"
    ]

    libreoffice = next(
        (p for p in possible_paths if Path(p).is_file()),
        None
    )

    # Fall back to PATH if neither standard location exists.
    if not libreoffice:
        libreoffice = (
            shutil.which("soffice.exe")
            or shutil.which("soffice")
            or shutil.which("libreoffice")
        )

    if not libreoffice:
        raise FileNotFoundError(
            "LibreOffice was not found. Checked standard installation "
            "paths and PATH. Please verify the location of soffice.exe."
        )

    word_path = Path(word_path).resolve()
    output_directory = Path(output_directory).resolve()
    output_directory.mkdir(parents=True, exist_ok=True)

    expected_pdf = output_directory / f"{word_path.stem}.pdf"

    print(f"Using LibreOffice: {libreoffice}")
    print("Converting Word report to PDF...")

    result = subprocess.run(
        [
            libreoffice,
            "--headless",
            "--convert-to", "pdf",
            "--outdir", str(output_directory),
            str(word_path)
        ],
        capture_output=True,
        text=True,
        timeout=120
    )

    if result.stdout:
        print(result.stdout)

    if result.stderr:
        print(result.stderr)

    if (
        result.returncode != 0
        or not expected_pdf.exists()
        or expected_pdf.stat().st_size == 0
    ):
        raise RuntimeError(
            "PDF conversion failed.\n"
            f"Return code: {result.returncode}\n"
            f"stdout: {result.stdout}\n"
            f"stderr: {result.stderr}"
        )

    print(f"PDF report saved: {expected_pdf}")
    return True



# ---------------------------------------------------------
# MAIN: GENERATE A SEPARATE WORD AND PDF FOR EACH WORKSHEET
# ---------------------------------------------------------
def main():
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_FILE}"
        )

    if not TEMPLATE_FILE.exists():
        raise FileNotFoundError(
            f"Word template not found: {TEMPLATE_FILE}"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("Reading Excel workbook...")
    districts = read_excel_data(EXCEL_FILE)

    if not districts:
        raise ValueError(
            "No district/location tables found. "
            "Check the workbook structure and Damaged headers."
        )

    print(f"Worksheets with data found: {len(districts)}")

    successful = []
    failed = []

    for district in districts:
        district_name = district["name"]

        print(f"\nProcessing worksheet: {district_name}")

        try:
            # Create one Word file from this worksheet.
            word_path = generate_word(
                TEMPLATE_FILE,
                OUTPUT_DIR,
                district
            )

            # Convert only that Word file to PDF.
            convert_to_pdf(word_path, OUTPUT_DIR)

            pdf_path = OUTPUT_DIR / f"{word_path.stem}.pdf"

            if not pdf_path.exists() or pdf_path.stat().st_size == 0:
                raise RuntimeError(
                    f"PDF was not created: {pdf_path}"
                )

            successful.append(district_name)
            print(f"Completed: {district_name}")

        except Exception as error:
            failed.append({
                "district": district_name,
                "error": str(error)
            })
            print(
                f"Failed for {district_name}: {error}"
            )

    print("\n========== AUTOMATION SUMMARY ==========")
    print(f"Successful worksheets: {len(successful)}")
    print(f"Failed worksheets: {len(failed)}")

    if failed:
        for item in failed:
            print(
                f"- {item['district']}: {item['error']}"
            )

        raise RuntimeError(
            f"{len(failed)} worksheet(s) failed. "
            "See the errors above."
        )

    print(f"All reports saved in: {OUTPUT_DIR}")
    print("Automation finished successfully.")


if __name__ == "__main__":
    main()
